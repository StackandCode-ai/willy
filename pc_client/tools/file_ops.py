"""
File and folder operations for Willy on the PC: read, write, list, move, copy, rename,
zip, unzip, make a folder, and delete (to the Recycle Bin, so it can be undone).

Paths can be absolute, use ~ / %VARS%, start with a known folder name ("Downloads/report.pdf",
"desktop"), or be a bare file name that is looked up in the usual folders.
"""

import os
import shutil
import zipfile
from pathlib import Path
from typing import Any, Dict, Optional

HOME = Path.home()
READ_LIMIT = 12_000  # characters returned to the brain
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".log", ".py", ".js", ".ts", ".html", ".css", ".xml",
                 ".yaml", ".yml", ".ini", ".cfg", ".bat", ".ps1", ".sql", ".env", ".dart", ".kt", ".java"}


def _known_folders() -> Dict[str, Path]:
    folders = {
        "downloads": HOME / "Downloads", "download": HOME / "Downloads",
        "documents": HOME / "Documents", "docs": HOME / "Documents",
        "desktop": HOME / "Desktop", "pictures": HOME / "Pictures", "photos": HOME / "Pictures",
        "music": HOME / "Music", "videos": HOME / "Videos", "home": HOME,
        "screenshots": HOME / "Pictures" / "Screenshots",
    }
    onedrive = os.environ.get("OneDrive")
    if onedrive:  # OneDrive often redirects Desktop / Documents / Pictures
        for key, sub in (("desktop", "Desktop"), ("documents", "Documents"), ("pictures", "Pictures")):
            if not folders[key].exists() and (Path(onedrive) / sub).exists():
                folders[key] = Path(onedrive) / sub
    return folders


def resolve(raw: str, must_exist: bool = False) -> Path:
    text = os.path.expandvars(os.path.expanduser(str(raw or "").strip().strip('"')))
    if not text:
        raise ValueError("No path given.")
    path = Path(text)
    if path.is_absolute():
        if must_exist and not path.exists():
            raise FileNotFoundError(f"There's nothing at {path}.")
        return path
    parts = Path(text.replace("\\", "/")).parts
    known = _known_folders()
    if parts and parts[0].lower() in known:
        return known[parts[0].lower()].joinpath(*parts[1:])
    if must_exist:
        from pc_client.tools.file_relay import find_file

        found, error = find_file(text)
        if found is None:
            raise FileNotFoundError(error)
        return found
    return known["documents"] / path


def _display(path: Path) -> str:
    try:
        return "~\\" + str(path.relative_to(HOME))
    except ValueError:
        return str(path)


def _to_recycle_bin(path: Path) -> None:
    from win32com.shell import shell, shellcon

    flags = shellcon.FOF_ALLOWUNDO | shellcon.FOF_NOCONFIRMATION | shellcon.FOF_SILENT | shellcon.FOF_NOERRORUI
    result, aborted = shell.SHFileOperation((0, shellcon.FO_DELETE, str(path), None, flags, None, None))
    if result != 0 or aborted:
        raise OSError(f"Windows couldn't move it to the Recycle Bin (code {result}).")


def _read(path: Path) -> Dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        try:
            import docx  # python-docx
        except ImportError:
            return {"success": False, "error": "Reading Word files needs python-docx on the PC."}
        text = "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            return {"success": False, "error": "Reading PDFs needs pypdf on the PC."}
        text = "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages[:30])
    else:
        raw = path.read_bytes()[: READ_LIMIT * 4]
        if b"\x00" in raw[:4096] and suffix not in TEXT_SUFFIXES:
            return {"success": False, "error": f"{path.name} isn't a text file, so I can't read it out."}
        text = raw.decode("utf-8", errors="replace")
    clipped = len(text) > READ_LIMIT
    return {"success": True, "path": str(path), "content": text[:READ_LIMIT], "truncated": clipped,
            "message": f"Read {path.name}" + (" (first part)." if clipped else ".")}


def _write(path: Path, content: str, overwrite: bool) -> Dict[str, Any]:
    if path.exists() and not overwrite:
        return {"success": False, "error": f"{_display(path)} already exists. Say if it should be replaced."}
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".docx":
        try:
            import docx
        except ImportError:
            return {"success": False, "error": "Creating Word files needs python-docx on the PC."}
        document = docx.Document()
        for line in (content or "").split("\n"):
            document.add_paragraph(line)
        document.save(str(path))
    else:
        path.write_text(content or "", encoding="utf-8")
    return {"success": True, "path": str(path), "message": f"Saved {path.name} in {_display(path.parent)}."}


def manage_file(op: str, path: str = "", destination: str = "", content: str = "",
                overwrite: bool = False) -> Dict[str, Any]:
    op = (op or "").strip().lower()
    try:
        if op in ("write", "create"):
            return _write(resolve(path), content, overwrite)
        if op in ("mkdir", "create_folder"):
            target = resolve(path)
            target.mkdir(parents=True, exist_ok=True)
            return {"success": True, "path": str(target), "message": f"Made the folder {_display(target)}."}
        if op == "list":
            folder = resolve(path or "desktop")
            if not folder.is_dir():
                return {"success": False, "error": f"{_display(folder)} isn't a folder."}
            entries = sorted(folder.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True)[:40]
            items = [{"name": e.name, "folder": e.is_dir(), "size": e.stat().st_size if e.is_file() else None}
                     for e in entries]
            return {"success": True, "path": str(folder), "items": items,
                    "message": f"{_display(folder)} has {len(list(folder.iterdir()))} items (newest first)."}

        source = resolve(path, must_exist=True)
        if not source.exists():
            return {"success": False, "error": f"There's nothing at {_display(source)}."}
        if op == "read":
            if source.is_dir():
                return manage_file("list", str(source))
            return _read(source)
        if op in ("move", "copy"):
            dest = resolve(destination) if destination else None
            if dest is None:
                return {"success": False, "error": f"Where should I {op} {source.name}?"}
            if dest.is_dir() or not dest.suffix:
                dest.mkdir(parents=True, exist_ok=True)
                dest = dest / source.name
            if dest.exists() and not overwrite:
                return {"success": False, "error": f"{_display(dest)} already exists. Say if it should be replaced."}
            if op == "move":
                shutil.move(str(source), str(dest))
            elif source.is_dir():
                shutil.copytree(source, dest, dirs_exist_ok=overwrite)
            else:
                shutil.copy2(source, dest)
            verb = "Moved" if op == "move" else "Copied"
            return {"success": True, "path": str(dest), "message": f"{verb} {source.name} to {_display(dest.parent)}."}
        if op == "rename":
            new_name = Path(str(destination or "")).name
            if not new_name:
                return {"success": False, "error": "What should the new name be?"}
            dest = source.with_name(new_name)
            if dest.exists():
                return {"success": False, "error": f"{new_name} already exists there."}
            source.rename(dest)
            return {"success": True, "path": str(dest), "message": f"Renamed {source.name} to {new_name}."}
        if op == "zip":
            base = resolve(destination) if destination else source.with_suffix("")
            archive = Path(shutil.make_archive(str(base.with_suffix("")), "zip",
                                               root_dir=str(source.parent), base_dir=source.name))
            return {"success": True, "path": str(archive),
                    "message": f"Zipped {source.name} into {archive.name} in {_display(archive.parent)}."}
        if op == "unzip":
            dest = resolve(destination) if destination else source.with_suffix("")
            with zipfile.ZipFile(source) as zf:
                for member in zf.namelist():  # refuse entries that would land outside dest
                    target = (dest / member).resolve()
                    if not str(target).startswith(str(dest.resolve())):
                        return {"success": False, "error": f"{source.name} has unsafe paths; not extracting."}
                zf.extractall(dest)
            return {"success": True, "path": str(dest), "message": f"Extracted {source.name} to {_display(dest)}."}
        if op == "delete":
            _to_recycle_bin(source)
            return {"success": True, "message": f"Moved {source.name} to the Recycle Bin (you can restore it from there)."}
        return {"success": False, "error": f"Unknown file operation '{op}'."}
    except (OSError, ValueError, zipfile.BadZipFile) as e:
        return {"success": False, "error": str(e)}


def set_wallpaper(path: str) -> Dict[str, Any]:
    import ctypes

    try:
        image = resolve(path, must_exist=True)
    except (OSError, ValueError) as e:
        return {"success": False, "error": str(e)}
    if not image.is_file():  # Windows "accepts" a missing file and shows a black desktop
        return {"success": False, "error": f"There's no picture at {image}."}
    if image.suffix.lower() not in (".jpg", ".jpeg", ".png", ".bmp", ".gif", ".webp"):
        return {"success": False, "error": f"{image.name} isn't a picture."}
    SPI_SETDESKWALLPAPER, UPDATE = 20, 0x01 | 0x02
    if not ctypes.windll.user32.SystemParametersInfoW(SPI_SETDESKWALLPAPER, 0, str(image), UPDATE):
        return {"success": False, "error": "Windows didn't accept that picture as the wallpaper."}
    return {"success": True, "message": f"Wallpaper set to {image.name}."}


def app_volume(app: str, level: Optional[int] = None, mute: Optional[bool] = None) -> Dict[str, Any]:
    """Volume / mute for one app's audio (e.g. only the browser), via its audio sessions."""
    from pycaw.pycaw import AudioUtilities

    wanted = (app or "").lower().replace(".exe", "").strip()
    aliases = {"browser": ("chrome", "msedge", "firefox", "brave", "opera"), "edge": ("msedge",),
               "spotify": ("spotify",), "chrome": ("chrome",)}
    names = aliases.get(wanted, (wanted,))
    touched = []
    for session in AudioUtilities.GetAllSessions():
        proc = session.Process
        if proc is None:
            continue
        pname = proc.name().lower().replace(".exe", "")
        if not any(n and n in pname for n in names):
            continue
        volume = session.SimpleAudioVolume
        if mute is not None:
            volume.SetMute(1 if mute else 0, None)
        if level is not None:
            volume.SetMasterVolume(max(0, min(100, int(level))) / 100.0, None)
        touched.append(proc.name())
    if not touched:
        return {"success": False, "error": f"{app} isn't playing any sound right now."}
    what = "Muted" if mute else ("Unmuted" if mute is False else f"Set to {level}%:")
    return {"success": True, "apps": sorted(set(touched)),
            "message": f"{what} {', '.join(sorted(set(touched)))}" + ("" if level is None else ".")}


def list_dir(path: str = "") -> Dict[str, Any]:
    """A folder's entries for the file browser; an empty path lists the drives and home folders."""
    if not (path or "").strip():
        drives = []
        for letter in "CDEFGHIJKLMNOPQRSTUVWXYZ":
            root = f"{letter}:\\"
            if not os.path.exists(root):
                continue
            try:
                usage = shutil.disk_usage(root)
                drives.append({"name": root, "label": f"{letter}: drive",
                               "free_gb": round(usage.free / 1e9, 1), "total_gb": round(usage.total / 1e9, 1)})
            except OSError:
                drives.append({"name": root, "label": f"{letter}: drive", "free_gb": None, "total_gb": None})
        seen = set()
        places = []
        for key in ("desktop", "downloads", "documents", "pictures", "music", "videos"):
            folder = _known_folders()[key]
            if folder.exists() and folder not in seen:
                seen.add(folder)
                places.append({"name": folder.name, "path": str(folder), "folder": True, "size": None,
                               "modified": folder.stat().st_mtime})
        return {"success": True, "path": "", "parent": None, "entries": places, "drives": drives,
                "message": f"{len(drives)} drives."}
    try:
        folder = resolve(path, must_exist=True)
    except (OSError, ValueError) as e:
        return {"success": False, "error": str(e)}
    if not folder.is_dir():
        return {"success": False, "error": f"{folder} isn't a folder."}
    entries = []
    try:
        with os.scandir(folder) as it:
            for entry in it:
                if entry.name.lower() in ("desktop.ini", "thumbs.db") or entry.name.startswith("$"):
                    continue
                try:
                    is_dir = entry.is_dir()
                    st = entry.stat()
                except OSError:
                    continue
                entries.append({"name": entry.name, "path": entry.path, "folder": is_dir,
                                "size": None if is_dir else st.st_size, "modified": st.st_mtime})
    except PermissionError:
        return {"success": False, "error": f"Windows doesn't allow opening {folder}."}
    entries.sort(key=lambda e: (not e["folder"], e["name"].lower()))
    parent = str(folder.parent) if folder.parent != folder else ""
    return {"success": True, "path": str(folder), "parent": parent, "entries": entries[:1000],
            "truncated": len(entries) > 1000, "message": f"{len(entries)} items in {folder.name or folder}."}


def open_file(path: str) -> Dict[str, Any]:
    """Opens a file (default app) or folder (Explorer) on the PC."""
    try:
        target = resolve(path, must_exist=True)
    except (OSError, ValueError) as e:
        return {"success": False, "error": str(e)}
    os.startfile(str(target))  # noqa: S606 - the user's own file, opened on their own PC
    return {"success": True, "message": f"Opened {target.name or target} on your PC."}
